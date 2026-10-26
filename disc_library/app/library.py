"""Library queries: browsing, locations, history, stats and Vibe Sesh picks."""

from __future__ import annotations

import random
from datetime import date, datetime, timedelta, timezone

from .db import Database

SLOTS_PER_CHANGER = 400
LOCATION_TYPES = ("changer", "shelf", "player", "loaned", "unassigned")

SORTS = {
    "artist": "a.artist COLLATE NOCASE, a.year, a.title COLLATE NOCASE",
    "title": "a.title COLLATE NOCASE",
    "year": "a.year DESC, a.artist COLLATE NOCASE",
    "added": "a.created_at DESC, a.id DESC",
    "played": "last_played DESC NULLS LAST, a.artist COLLATE NOCASE",
    "value": "a.value DESC NULLS LAST, a.paid DESC NULLS LAST, a.artist COLLATE NOCASE",
}

VALUE_TIERS = ((9, "$"), (25, "$$"))  # under $9, $9-24.99, $25 and up


def value_tier(value: float | None) -> str | None:
    if value is None:
        return None
    for limit, mark in VALUE_TIERS:
        if value < limit:
            return mark
    return "$$$"


def format_group(fmt: str | None, media_kind: str) -> str:
    f = (fmt or "").lower()
    if media_kind == "video":
        return "Blu-ray" if "blu" in f else "DVD" if "dvd" in f else "Video"
    if "blu" in f:
        return "Blu-ray Audio"
    if "sacd" in f:
        return "SACD"
    if "dvd" in f:
        return "DVD-Audio"
    return "CD"


class RemoveError(ValueError):
    pass

ALBUM_SELECT = """
SELECT a.*, d.id AS disc_id, d.location_type, d.changer_id, d.slot, d.location_note,
       (SELECT MAX(started_at) FROM sessions s WHERE s.album_id = a.id) AS last_played,
       (SELECT COUNT(*) FROM sessions s WHERE s.album_id = a.id AND s.completed = 1) AS full_listens,
       (SELECT COUNT(*) FROM discs dd WHERE dd.album_id = a.id) AS disc_count,
       (SELECT COUNT(*) FROM discs dd WHERE dd.album_id = a.id AND dd.location_type = 'changer') AS discs_in_changer
FROM albums a
LEFT JOIN discs d ON d.id = (SELECT id FROM discs WHERE album_id = a.id ORDER BY disc_number LIMIT 1)
"""

ALBUM_YEAR = "COALESCE(a.original_year, a.year)"


class LocationError(ValueError):
    pass


class Library:
    def __init__(self, db: Database, changer_ids: list[int]):
        self.db = db
        self.changer_ids = changer_ids

    # --- browse ----------------------------------------------------------------
    def albums(self, q: str | None = None, genre: str | None = None, favorites: bool = False,
               media_kind: str | None = "audio", status: str | None = "owned", sort: str = "artist",
               limit: int = 1000, offset: int = 0, decade: int | None = None) -> list[dict]:
        where, params = ["a.removed_at IS NULL"], []
        if decade:
            where.append(f"{ALBUM_YEAR} BETWEEN ? AND ?")
            params += [decade, decade + 9]
        if q:
            where.append("""(a.artist LIKE ? OR a.title LIKE ? OR a.label LIKE ?
                             OR EXISTS (SELECT 1 FROM tracks t WHERE t.album_id = a.id AND t.title LIKE ?))""")
            params += [f"%{q}%"] * 4
        if genre:
            where.append("(a.genre = ? OR a.subgenre = ?)")
            params += [genre, genre]
        if favorites:
            where.append("a.favorite = 1")
        if media_kind:
            where.append("a.media_kind = ?")
            params.append(media_kind)
        if status:
            where.append("a.status = ?")
            params.append(status)
        sql = ALBUM_SELECT + (" WHERE " + " AND ".join(where) if where else "")
        sql += f" ORDER BY {SORTS.get(sort, SORTS['artist'])} LIMIT ? OFFSET ?"
        return self.db.query(sql, (*params, limit, offset))

    def album(self, album_id: int) -> dict | None:
        a = self.db.one(ALBUM_SELECT + " WHERE a.id = ?", (album_id,))
        if not a:
            return None
        a["discs"] = self.db.query("SELECT * FROM discs WHERE album_id = ? ORDER BY disc_number", (album_id,))
        a["tracks"] = self.db.query(
            "SELECT disc_number, number, title, duration_sec FROM tracks WHERE album_id = ? ORDER BY disc_number, number",
            (album_id,),
        )
        a["resume"] = self.db.one(
            """SELECT last_track, ended_at FROM sessions WHERE album_id = ? AND completed = 0
               AND ended_at IS NOT NULL ORDER BY started_at DESC, id DESC LIMIT 1""",
            (album_id,),
        )
        return a

    def genres(self) -> list[dict]:
        return self.db.query(
            """SELECT genre AS name, COUNT(*) AS albums FROM albums
               WHERE genre IS NOT NULL AND media_kind = 'audio' AND status = 'owned' AND removed_at IS NULL
               GROUP BY genre ORDER BY albums DESC"""
        )

    def decades(self) -> list[dict]:
        return self.db.query(
            f"""SELECT ({ALBUM_YEAR} / 10) * 10 AS decade, COUNT(*) AS albums FROM albums a
                WHERE {ALBUM_YEAR} IS NOT NULL AND a.media_kind = 'audio' AND a.status = 'owned'
                  AND a.removed_at IS NULL GROUP BY 1 ORDER BY 1""")

    def set_favorite(self, album_id: int, favorite: bool) -> None:
        self.db.execute("UPDATE albums SET favorite = ? WHERE id = ?", (int(favorite), album_id))

    def set_tags(self, album_id: int, genre: str | None, subgenre: str | None, mood_tags: str | None) -> None:
        self.db.execute(
            "UPDATE albums SET genre = ?, subgenre = ?, mood_tags = ?, updated_at = datetime('now') WHERE id = ?",
            (genre, subgenre, mood_tags, album_id),
        )

    # --- locations -------------------------------------------------------------
    def set_location(self, disc_id: int, location_type: str, changer_id: int | None = None,
                     slot: int | None = None, note: str | None = None) -> dict:
        if location_type not in LOCATION_TYPES:
            raise LocationError(f"Unknown location: {location_type}")
        if location_type == "changer":
            if changer_id not in self.changer_ids:
                raise LocationError("Pick one of the configured changers.")
            if not slot or not 1 <= slot <= SLOTS_PER_CHANGER:
                raise LocationError(f"Slot must be 1-{SLOTS_PER_CHANGER}.")
            taken = self.db.one(
                """SELECT a.artist, a.title FROM discs d JOIN albums a ON a.id = d.album_id
                   WHERE d.changer_id = ? AND d.slot = ? AND d.id != ?""",
                (changer_id, slot, disc_id),
            )
            if taken:
                raise LocationError(f"Slot {slot} already holds {taken['artist']} - {taken['title']}.")
        else:
            changer_id = slot = None
        self.db.execute(
            "UPDATE discs SET location_type = ?, changer_id = ?, slot = ?, location_note = ? WHERE id = ?",
            (location_type, changer_id, slot, note, disc_id),
        )
        return self.db.one("SELECT * FROM discs WHERE id = ?", (disc_id,))

    def free_slot(self, changer_id: int | None = None) -> dict | None:
        """First empty slot, on the given changer or on whichever has the most room."""
        candidates = [changer_id] if changer_id else sorted(
            self.changer_ids, key=lambda c: self._used(c)
        )
        for cid in candidates:
            used = {r["slot"] for r in self.db.query("SELECT slot FROM discs WHERE changer_id = ?", (cid,))}
            for s in range(1, SLOTS_PER_CHANGER + 1):
                if s not in used:
                    return {"changer_id": cid, "slot": s}
        return None

    def _used(self, changer_id: int) -> int:
        return self.db.one("SELECT COUNT(*) AS n FROM discs WHERE changer_id = ?", (changer_id,))["n"]

    def slot_map(self) -> list[dict]:
        return [
            {"changer_id": cid, "used": self._used(cid), "capacity": SLOTS_PER_CHANGER}
            for cid in self.changer_ids
        ]

    # --- removing ----------------------------------------------------------------
    def remove_preview(self, album_id: int) -> dict:
        a = self.db.one("SELECT id, artist, title, removed_at FROM albums WHERE id = ?", (album_id,))
        if not a:
            raise RemoveError("Album not found.")
        discs = self.db.query("SELECT disc_number, location_type, changer_id, slot FROM discs WHERE album_id = ? "
                              "ORDER BY disc_number", (album_id,))
        listens = self.db.one("SELECT COUNT(*) AS n FROM sessions WHERE album_id = ?", (album_id,))["n"]
        return {**a, "discs": discs, "listens": listens,
                "slots": [d for d in discs if d["location_type"] in ("changer", "to_pull")]}

    # --- linking discs into a set by hand --------------------------------------------
    def link_disc(self, album_id: int, other_id: int, as_disc: int) -> dict:
        """Fold a one-disc album into this album as disc `as_disc` (for sets the spreadsheet names inconsistently)."""
        if album_id == other_id:
            raise ValueError("Pick a different album.")
        a = self.db.one("SELECT * FROM albums WHERE id = ? AND removed_at IS NULL", (album_id,))
        o = self.db.one("SELECT * FROM albums WHERE id = ? AND removed_at IS NULL", (other_id,))
        if not a or not o:
            raise ValueError("Album not found.")
        if not 1 <= as_disc <= 20:
            raise ValueError("Disc number must be 1-20.")
        odiscs = self.db.query("SELECT * FROM discs WHERE album_id = ?", (other_id,))
        if len(odiscs) != 1:
            raise ValueError(f"{o['title']} already has several discs; link its discs the other way round.")
        if self.db.one("SELECT 1 FROM discs WHERE album_id = ? AND disc_number = ?", (album_id, as_disc)):
            raise ValueError(f"{a['title']} already has a disc {as_disc}.")
        with self.db.connect() as con:
            con.execute("UPDATE discs SET album_id = ?, disc_number = ? WHERE id = ?", (album_id, as_disc, odiscs[0]["id"]))
            con.execute("UPDATE sessions SET album_id = ? WHERE album_id = ?", (album_id, other_id))
            con.execute("INSERT OR REPLACE INTO import_aliases (import_key, album_id, disc_number) VALUES (?, ?, ?)",
                        (o["import_key"], album_id, as_disc))
            con.execute("UPDATE import_aliases SET album_id = ? WHERE album_id = ?", (album_id, other_id))
            con.execute("DELETE FROM albums WHERE id = ?", (other_id,))
            con.execute("""UPDATE albums SET paid = COALESCE(paid, ?), paid_source = COALESCE(paid_source, ?),
                           enrich_status = CASE WHEN pressing_locked = 1 THEN enrich_status ELSE 'pending' END
                           WHERE id = ?""", (o["paid"], o["paid_source"] if o["paid"] is not None else None, album_id))
        return self.album(album_id)

    def unlink_disc(self, disc_id: int) -> dict:
        """Split one disc of a set back out into its own album."""
        d = self.db.one("SELECT * FROM discs WHERE id = ?", (disc_id,))
        if not d:
            raise ValueError("Disc not found.")
        a = self.db.one("SELECT * FROM albums WHERE id = ?", (d["album_id"],))
        if self.db.one("SELECT COUNT(*) AS n FROM discs WHERE album_id = ?", (a["id"],))["n"] < 2:
            raise ValueError("This album has only one disc.")
        alias = self.db.one("SELECT import_key FROM import_aliases WHERE album_id = ? AND disc_number = ?",
                            (a["id"], d["disc_number"]))
        key = alias["import_key"] if alias else f"split|{a['import_key']}|{d['disc_number']}"
        with self.db.connect() as con:
            cur = con.execute(
                """INSERT INTO albums (import_key, artist, title, year, label, format, media_kind, status, source_sheet)
                   VALUES (?,?,?,?,?,?,?,?,?)""",
                (key, a["artist"], f"{a['title']} — Disc {d['disc_number']}", a["year"], a["label"], a["format"],
                 a["media_kind"], a["status"], a["source_sheet"]))
            con.execute("UPDATE discs SET album_id = ?, disc_number = 1 WHERE id = ?", (cur.lastrowid, disc_id))
            con.execute("DELETE FROM import_aliases WHERE import_key = ?", (key,))
            con.execute("DELETE FROM tracks WHERE album_id = ? AND disc_number = ?", (a["id"], d["disc_number"]))
        return self.album(cur.lastrowid)

    def remove(self, album_id: int) -> dict:
        info = self.remove_preview(album_id)
        if info["removed_at"]:
            return info
        with self.db.connect() as con:
            # Discs in a changer keep their slot, marked "to pull", so no new disc is sent into a slot
            # that is still physically full. The slot frees when the disc is marked pulled.
            con.execute("UPDATE discs SET location_type = 'to_pull' WHERE album_id = ? AND location_type = 'changer'",
                        (album_id,))
            con.execute("UPDATE albums SET removed_at = datetime('now') WHERE id = ?", (album_id,))
        return self.remove_preview(album_id)

    def to_pull(self) -> list[dict]:
        return self.db.query(
            """SELECT d.id AS disc_id, d.disc_number, d.changer_id, d.slot, a.id AS album_id, a.artist, a.title,
                      a.art_url, (SELECT COUNT(*) FROM discs x WHERE x.album_id = a.id) AS disc_count
               FROM discs d JOIN albums a ON a.id = d.album_id
               WHERE d.location_type = 'to_pull' ORDER BY d.changer_id, d.slot""")

    def mark_pulled(self, disc_ids: list[int]) -> int:
        n = 0
        with self.db.connect() as con:
            for disc_id in disc_ids:
                cur = con.execute(
                    """UPDATE discs SET prev_changer_id = changer_id, prev_slot = slot, changer_id = NULL, slot = NULL,
                       location_type = 'unassigned' WHERE id = ? AND location_type = 'to_pull'""", (disc_id,))
                n += cur.rowcount
        return n

    def restore(self, album_id: int) -> dict:
        """Bring a removed album back, into its old slot when that slot is still free."""
        a = self.db.one("SELECT id, removed_at FROM albums WHERE id = ?", (album_id,))
        if not a:
            raise RemoveError("Album not found.")
        back = []
        with self.db.connect() as con:
            con.execute("UPDATE albums SET removed_at = NULL WHERE id = ?", (album_id,))
            for d in con.execute("SELECT * FROM discs WHERE album_id = ?", (album_id,)).fetchall():
                if d["location_type"] == "to_pull":  # never pulled: it's still sitting in its slot
                    con.execute("UPDATE discs SET location_type = 'changer' WHERE id = ?", (d["id"],))
                    back.append({"changer_id": d["changer_id"], "slot": d["slot"]})
                    continue
                if d["prev_slot"] and d["prev_changer_id"] in self.changer_ids:
                    taken = con.execute("SELECT 1 FROM discs WHERE changer_id = ? AND slot = ?",
                                        (d["prev_changer_id"], d["prev_slot"])).fetchone()
                    if not taken:
                        con.execute("""UPDATE discs SET location_type = 'changer', changer_id = ?, slot = ?,
                                       prev_changer_id = NULL, prev_slot = NULL WHERE id = ?""",
                                    (d["prev_changer_id"], d["prev_slot"], d["id"]))
                        back.append({"changer_id": d["prev_changer_id"], "slot": d["prev_slot"]})
        return {"id": album_id, "back_in_slot": back}

    def removed(self) -> list[dict]:
        return self.db.query(
            """SELECT a.id, a.artist, a.title, a.format, a.art_url, a.removed_at FROM albums a
               WHERE a.removed_at IS NOT NULL ORDER BY a.removed_at DESC""")

    # --- money ---------------------------------------------------------------------
    def set_money(self, album_id: int, changes: dict) -> dict:
        for key in ("paid", "value"):
            if key not in changes:
                continue
            v = changes[key]
            if v in ("", None):
                v = None
            else:
                try:
                    v = round(float(str(v).replace("$", "").replace(",", "")), 2)
                except ValueError as exc:
                    raise ValueError(f"{key.capitalize()} must be a dollar amount.") from exc
                if v < 0:
                    raise ValueError(f"{key.capitalize()} can't be negative.")
            src = "manual" if v is not None else None
            if key == "paid":
                self.db.execute("UPDATE albums SET paid = ?, paid_source = ? WHERE id = ?", (v, src, album_id))
            else:
                self.db.execute("UPDATE albums SET value = ?, value_source = ?, value_at = datetime('now') WHERE id = ?",
                                (v, src, album_id))
        return self.db.one("SELECT paid, paid_source, value, value_source, value_at FROM albums WHERE id = ?", (album_id,))

    # --- collection --------------------------------------------------------------
    def collection(self, changer_names: dict[int, str] | None = None) -> dict:
        albums = self.db.query(
            """SELECT a.id, a.artist, a.title, a.format, a.media_kind, a.status, a.runtime_sec, a.paid, a.value,
                      a.art_url, (SELECT COUNT(*) FROM sessions s WHERE s.album_id = a.id AND s.completed = 1) AS full_listens,
                      (SELECT COUNT(*) FROM discs d WHERE d.album_id = a.id) AS discs
               FROM albums a WHERE a.removed_at IS NULL""")
        owned = [a for a in albums if a["status"] == "owned"]
        formats: dict[str, int] = {}
        for a in owned:
            g = format_group(a["format"], a["media_kind"])
            formats[g] = formats.get(g, 0) + 1
        disc_rows = self.db.query(
            """SELECT d.location_type, d.changer_id FROM discs d JOIN albums a ON a.id = d.album_id
               WHERE a.removed_at IS NULL AND a.status = 'owned'""")
        names = changer_names or {}
        changers = [{"id": cid, "name": names.get(cid, f"Changer {cid}"), "used": self._used(cid),
                     "capacity": SLOTS_PER_CHANGER} for cid in self.changer_ids]
        audio = [a for a in owned if a["media_kind"] == "audio"]
        timed = [a for a in audio if a["runtime_sec"]]
        paid = [a for a in owned if a["paid"] is not None]
        valued = [a for a in owned if a["value"] is not None]
        tiers = {"$": 0, "$$": 0, "$$$": 0}
        for a in valued:
            tiers[value_tier(a["value"])] += 1
        played_paid = [a for a in paid if a["full_listens"]]
        listens = sum(a["full_listens"] for a in played_paid)
        top = sorted(valued, key=lambda a: a["value"], reverse=True)[:10]
        return {
            "albums": len(owned),
            "on_order": sum(1 for a in albums if a["status"] == "on_order"),
            "discs": len(disc_rows),
            "audio": len(audio),
            "video": sum(1 for a in owned if a["media_kind"] == "video"),
            "formats": dict(sorted(formats.items(), key=lambda kv: -kv[1])),
            "changers": changers,
            "not_in_changer": sum(1 for d in disc_rows if d["location_type"] != "changer"),
            "runtime_sec": sum(a["runtime_sec"] for a in timed),
            "runtime_known": len(timed),
            "money": {
                "paid_total": round(sum(a["paid"] for a in paid), 2),
                "paid_known": len(paid),
                "value_total": round(sum(a["value"] for a in valued), 2),
                "value_known": len(valued),
                "value_avg": round(sum(a["value"] for a in valued) / len(valued), 2) if valued else None,
                "tiers": tiers,
                "cost_per_listen": round(sum(a["paid"] for a in played_paid) / listens, 2) if listens else None,
                "top": [{"id": a["id"], "artist": a["artist"], "title": a["title"], "value": a["value"],
                         "paid": a["paid"], "art_url": a["art_url"]} for a in top],
            },
        }

    # --- history and stats -----------------------------------------------------
    def history(self, limit: int = 50) -> list[dict]:
        return self.db.query(
            """SELECT s.*, a.artist, a.title, a.year FROM sessions s JOIN albums a ON a.id = s.album_id
               WHERE a.removed_at IS NULL ORDER BY s.started_at DESC, s.id DESC LIMIT ?""",
            (limit,),
        )

    def continue_listening(self, limit: int = 5) -> list[dict]:
        """Albums whose most recent session stopped partway."""
        return self.db.query(
            """SELECT a.id AS album_id, a.artist, a.title, s.last_track, s.ended_at,
                      (SELECT COUNT(*) FROM tracks t WHERE t.album_id = a.id AND t.disc_number = 1) AS track_count
               FROM sessions s JOIN albums a ON a.id = s.album_id
               WHERE a.removed_at IS NULL AND s.id = (SELECT id FROM sessions s2 WHERE s2.album_id = s.album_id
                             ORDER BY s2.started_at DESC, s2.id DESC LIMIT 1)
                 AND s.completed = 0 AND s.ended_at IS NOT NULL AND s.end_reason != 'finished'
               ORDER BY s.ended_at DESC LIMIT ?""",
            (limit,),
        )

    def stats(self, since: str | None = None) -> dict:
        """since: ISO date (inclusive) or None for all time."""
        cond, params = ("WHERE a.removed_at IS NULL AND s.started_at >= ?", (since,)) if since \
            else ("WHERE a.removed_at IS NULL", ())
        totals = self.db.one(
            f"""SELECT COALESCE(SUM(s.completed), 0) AS full_listens,
                       COALESCE(SUM(CASE WHEN s.ended_at IS NOT NULL AND s.completed = 0 THEN 1 ELSE 0 END), 0)
                           AS partial_listens,
                       COALESCE(SUM(s.played_sec), 0) AS played_sec,
                       COUNT(DISTINCT s.album_id) AS albums_played
                FROM sessions s JOIN albums a ON a.id = s.album_id {cond}""",
            params,
        )
        owned = self.db.one(
            "SELECT COUNT(*) AS n FROM albums WHERE status = 'owned' AND media_kind = 'audio' AND removed_at IS NULL"
        )["n"]
        by_genre = self.db.query(
            f"""SELECT COALESCE(a.genre, 'Untagged') AS genre, SUM(s.completed) AS full_listens
                FROM sessions s JOIN albums a ON a.id = s.album_id {cond}
                GROUP BY 1 HAVING SUM(s.completed) > 0 ORDER BY 2 DESC""",
            params,
        )
        most_played = self.db.query(
            f"""SELECT a.id AS album_id, a.artist, a.title, COUNT(*) AS plays, SUM(s.completed) AS full_listens
                FROM sessions s JOIN albums a ON a.id = s.album_id {cond}
                GROUP BY a.id ORDER BY plays DESC, full_listens DESC LIMIT 10""",
            params,
        )
        days = [
            date.fromisoformat(r["d"])
            for r in self.db.query("SELECT DISTINCT date(started_at) AS d FROM sessions ORDER BY d")
        ]
        current, longest = _streaks(days, datetime.now(timezone.utc).date())  # session times are stored in UTC
        return {
            "full_listens": totals["full_listens"],
            "partial_listens": totals["partial_listens"],
            "hours": round(totals["played_sec"] / 3600, 1),
            "albums_played": totals["albums_played"],
            "albums_owned": owned,
            "streak_days": current,
            "longest_streak_days": longest,
            "by_genre": by_genre,
            "most_played": most_played,
        }

    # --- pressings ---------------------------------------------------------------
    PRESSING_FILTERS = {
        "review": """a.pressing_locked = 0 AND a.enrich_status IN ('matched', 'manual')
                     AND COALESCE(a.match_method, 'title') != 'catalog'""",
        "catalog": "a.pressing_locked = 0 AND a.enrich_status = 'matched' AND a.match_method = 'catalog'",
        "locked": "a.pressing_locked = 1",
        "unmatched": "a.pressing_locked = 0 AND a.enrich_status IN ('pending', 'not_found', 'error')",
        "skipped": "a.enrich_status = 'skipped'",
    }

    def pressing_counts(self) -> dict:
        base = "FROM albums a WHERE a.media_kind = 'audio' AND a.status = 'owned' AND a.removed_at IS NULL AND "
        return {k: self.db.one(f"SELECT COUNT(*) AS n {base} {w}")["n"] for k, w in self.PRESSING_FILTERS.items()}

    def pressings(self, which: str = "review", limit: int = 1000) -> list[dict]:
        where = self.PRESSING_FILTERS.get(which, self.PRESSING_FILTERS["review"])
        rows = self.db.query(
            f"""SELECT a.id, a.artist, a.title, a.year, a.format, a.catalog_no, a.catalog_verified, a.barcode,
                       a.pressing_locked, a.match_method, a.enrich_status, a.enrich_note, a.art_url,
                       a.rel_label, a.rel_catno, a.rel_barcode, a.rel_country, a.rel_date, a.rel_format,
                       a.rel_track_count, a.mbid_release
                FROM albums a WHERE a.media_kind = 'audio' AND a.status = 'owned' AND a.removed_at IS NULL AND {where}
                ORDER BY a.artist COLLATE NOCASE, a.year, a.title COLLATE NOCASE LIMIT ?""",
            (limit,),
        )
        for r in rows:
            mine = r["catalog_verified"] or r["catalog_no"]
            r["catalog_mismatch"] = bool(mine and r["rel_catno"] and not (
                set(_cat_forms(mine)) & set(_cat_forms(r["rel_catno"]))))
        return rows

    def export_pressings(self, path) -> int:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Pressings"
        cols = ["Artist", "Album", "Your catalog no.", "Verified catalog no.", "Barcode", "Locked",
                "Match", "Label", "Country", "Release date", "Format", "Tracks", "MusicBrainz release"]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="FF2C2926")
        rows = self.db.query(
            """SELECT * FROM albums WHERE media_kind = 'audio' AND status = 'owned' AND removed_at IS NULL
               ORDER BY artist COLLATE NOCASE, year, title COLLATE NOCASE""")
        for r in rows:
            ws.append([r["artist"], r["title"], r["catalog_no"], r["catalog_verified"], r["barcode"] or r["rel_barcode"],
                       "Yes" if r["pressing_locked"] else "", r["match_method"] or r["enrich_status"],
                       r["rel_label"], r["rel_country"], r["rel_date"], r["rel_format"], r["rel_track_count"],
                       f"https://musicbrainz.org/release/{r['mbid_release']}" if r["mbid_release"] else None])
        for col, width in zip("ABCDEFGHIJKLM", (24, 34, 18, 18, 16, 8, 10, 22, 8, 12, 12, 7, 60)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        wb.save(path)
        return len(rows)

    # --- Vibe Sesh -------------------------------------------------------------
    def vibe(self, mood: str | None = None, energy: str | None = None, time_budget: str | None = None,
             familiarity: str | None = None, count: int = 3, seed: int | None = None) -> list[dict]:
        """Three playable picks, each with a reason built only from known facts.

        Mood and energy match against an album's mood tags; until the genre
        and mood tagging pass is done, those filters are skipped rather than
        guessed, and the picks lean on time and familiarity alone.
        """
        rows = self.albums(limit=5000)
        rows = [r for r in rows if r["location_type"] == "changer"]
        rng = random.Random(seed)
        cutoff = (datetime.now(timezone.utc) - timedelta(days=90)).strftime("%Y-%m-%d")
        max_runtime = {"under_45": 45 * 60, "about_hour": 70 * 60}.get(time_budget or "")

        scored = []
        for r in rows:
            reasons, score = [], rng.random()
            tags = {t.strip().lower() for t in (r["mood_tags"] or "").split(",") if t.strip()}
            for want in (mood, energy):
                if want and tags:
                    if want.lower() in tags:
                        score += 2
                        reasons.append(want.capitalize())
                    else:
                        score -= 1
            if max_runtime and r["runtime_sec"]:
                if r["runtime_sec"] <= max_runtime:
                    score += 1
                    reasons.append(f"{round(r['runtime_sec'] / 60)} minutes")
                else:
                    continue
            if familiarity == "rediscover":
                if not r["last_played"]:
                    score += 1.5
                    reasons.append("Not played here yet")
                elif r["last_played"] < cutoff:
                    score += 1
                    reasons.append(f"Last played {r['last_played'][:10]}")
                else:
                    score -= 2
            elif familiarity == "favorite":
                if r["favorite"]:
                    score += 2
                    reasons.append("A favorite")
                if r["full_listens"]:
                    score += min(r["full_listens"], 5) * 0.3
                    reasons.append(f"{r['full_listens']} full listens")
            if r["genre"]:
                reasons.append(r["genre"])
            scored.append((score, r, reasons))

        scored.sort(key=lambda x: x[0], reverse=True)
        return [
            {**r, "reason": " · ".join(reasons) if reasons else "In a changer and ready to play"}
            for _, r, reasons in scored[:count]
        ]


def _cat_forms(s: str) -> list[str]:
    from .enrich import catalog_candidates
    return [c for part in s.split(",") for c in catalog_candidates(part)]


def _streaks(days: list[date], today: date) -> tuple[int, int]:
    if not days:
        return 0, 0
    longest = run = 1
    for prev, cur in zip(days, days[1:]):
        run = run + 1 if (cur - prev).days == 1 else 1
        longest = max(longest, run)
    current = 0
    if days[-1] in (today, today - timedelta(days=1)):
        current = 1
        for prev, cur in zip(reversed(days[:-1]), reversed(days[1:])):
            if (cur - prev).days == 1:
                current += 1
            else:
                break
    return current, longest
