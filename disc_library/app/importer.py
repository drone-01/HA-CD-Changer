"""Import CD_Collection.xlsx into the library.

Matches columns by header name, so it copes with the three tabs (CD
Collection, Non-CD Formats, Video) and with columns being reordered.
Re-running it is safe: catalog fields are refreshed, while locations,
genres, favorites, art and history are left alone.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .db import Database
from .sets import merge_sets, set_key, set_member

# header text (lower-case, trimmed) -> field
HEADER_MAP = {
    "artist": "artist",
    "album": "title",
    "title": "title",
    "format": "format",
    "year": "year",
    "label": "label",
    "status": "status",
    "edition / notes": "notes",
    "notes": "notes",
    "edition": "notes",
}

STATUS_MAP = {
    "in collection": "owned",
    "owned": "owned",
    "on order": "on_order",
}


def _field_for(header: str) -> str | None:
    h = header.strip().lower()
    if "catalog" in h:
        return "catalog_no"
    return HEADER_MAP.get(h)


def _year(value) -> int | None:
    if value is None:
        return None
    m = re.search(r"(19|20)\d{2}", str(value))
    return int(m.group(0)) if m else None


def _clean(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


PRICE_TOTAL = re.compile(r"order total(?: of)?\s*(?:US\s*)?\$\s?(\d[\d,]*(?:\.\d{1,2})?)", re.I)
PRICE_FOR = re.compile(r"\bfor\s+(?:US\s*)?\$\s?(\d[\d,]*(?:\.\d{1,2})?)", re.I)


def paid_from_notes(notes: str | None) -> float | None:
    """The price in the eBay notes: the order total when given, else the "bought for $X" price."""
    if not notes:
        return None
    m = PRICE_TOTAL.search(notes) or PRICE_FOR.search(notes)
    if not m:
        return None
    try:
        return round(float(m.group(1).replace(",", "")), 2)
    except ValueError:
        return None


def import_key(artist: str, title: str, fmt: str, catalog: str | None) -> str:
    norm = lambda s: re.sub(r"\s+", " ", (s or "").strip().lower())
    return "|".join([norm(artist), norm(title), norm(fmt), norm(catalog)])


def parse_workbook(path: Path | str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    for ws in wb.worksheets:
        header_idx = None
        columns: dict[int, str] = {}
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if header_idx is None:
                labels = [str(v).strip().lower() if v is not None else "" for v in row]
                if "artist" in labels:
                    header_idx = r_idx
                    for c_idx, v in enumerate(row):
                        if v is None:
                            continue
                        f = _field_for(str(v))
                        if f and f not in columns.values():
                            columns[c_idx] = f
                continue
            rec = {f: row[c] if c < len(row) else None for c, f in columns.items()}
            artist, title = _clean(rec.get("artist")), _clean(rec.get("title"))
            if not artist or not title:
                continue
            fmt = _clean(rec.get("format"))
            is_video = "video" in ws.title.lower() or bool(fmt and re.search(r"video|\bdvd\b", fmt, re.I))
            fmt = fmt or ("DVD/Blu-ray" if is_video else "CD")
            status_raw = (_clean(rec.get("status")) or "").lower()
            rows.append(
                {
                    "artist": artist,
                    "title": title,
                    "format": fmt,
                    "media_kind": "video" if is_video else "audio",
                    "catalog_no": _clean(rec.get("catalog_no")),
                    "year": _year(rec.get("year")),
                    "label": _clean(rec.get("label")),
                    "edition_notes": _clean(rec.get("notes")),
                    "status": STATUS_MAP.get(status_raw, "owned"),
                    "paid": paid_from_notes(_clean(rec.get("notes"))),
                    "source_sheet": ws.title,
                }
            )
    wb.close()
    return rows


def import_workbook(db: Database, path: Path | str) -> dict:
    """Load the spreadsheet. Each disc of a multi-disc set becomes a disc of one album."""
    rows = parse_workbook(path)
    merge_sets(db)  # fold any one-row-per-disc albums from before sets existed
    by_sheet: dict[str, int] = {}
    groups: dict[str, list[tuple[int, dict]]] = {}
    order: list[str] = []
    for r in rows:
        by_sheet[r["source_sheet"]] = by_sheet.get(r["source_sheet"], 0) + 1
        member = set_member(r["title"], r["edition_notes"])
        if member:
            base, n = member
            key = set_key(r["artist"], base, r["format"])
            r = {**r, "title": base}
        else:
            key, n = import_key(r["artist"], r["title"], r["format"], r["catalog_no"]), 1
        if key not in groups:
            order.append(key)
            groups[key] = []
        groups[key].append((n, r))

    added = updated = removed = 0
    with db.connect() as con:
        for key in order:
            members = sorted(groups[key], key=lambda m: m[0])
            r = members[0][1]  # the lowest disc speaks for the set
            numbers = sorted({n for n, _ in members})
            paid = next((m["paid"] for _, m in members if m["paid"] is not None), None)
            catalog = next((m["catalog_no"] for _, m in members if m["catalog_no"]), None)
            status = "owned" if any(m["status"] == "owned" for _, m in members) else r["status"]
            alias = con.execute("SELECT album_id, disc_number FROM import_aliases WHERE import_key = ?", (key,)).fetchone()
            if alias:  # linked by hand into another album as one of its discs
                con.execute("INSERT OR IGNORE INTO discs (album_id, disc_number) VALUES (?, ?)",
                            (alias["album_id"], alias["disc_number"]))
                updated += 1
                continue
            existing = con.execute("SELECT id, removed_at FROM albums WHERE import_key = ?", (key,)).fetchone()
            if existing and existing["removed_at"]:
                removed += 1  # removed in the app: stays removed until restored there
                continue
            if existing:
                album_id = existing["id"]
                con.execute(
                    """UPDATE albums SET artist=?, title=?, year=?, label=?, catalog_no=?, format=?,
                       media_kind=?, edition_notes=?, status=?, source_sheet=?, updated_at=datetime('now')
                       WHERE id=?""",
                    (r["artist"], r["title"], r["year"], r["label"], catalog, r["format"],
                     r["media_kind"], r["edition_notes"], status, r["source_sheet"], album_id),
                )
                con.execute(
                    """UPDATE albums SET paid = ?, paid_source = CASE WHEN ? IS NULL THEN NULL ELSE 'notes' END
                       WHERE id = ? AND COALESCE(paid_source, 'notes') = 'notes'""",
                    (paid, paid, album_id))
                updated += 1
            else:
                cur = con.execute(
                    """INSERT INTO albums (import_key, artist, title, year, label, catalog_no, format,
                       media_kind, edition_notes, status, source_sheet, paid, paid_source)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (key, r["artist"], r["title"], r["year"], r["label"], catalog, r["format"],
                     r["media_kind"], r["edition_notes"], status, r["source_sheet"],
                     paid, "notes" if paid is not None else None),
                )
                album_id = cur.lastrowid
                added += 1
            for n in numbers:
                con.execute("INSERT OR IGNORE INTO discs (album_id, disc_number) VALUES (?, ?)", (album_id, n))
    return {"rows": len(rows), "added": added, "updated": updated, "skipped_removed": removed,
            "by_sheet": by_sheet}
