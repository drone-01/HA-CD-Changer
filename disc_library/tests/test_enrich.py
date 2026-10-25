import asyncio

import httpx
import pytest

from app.enrich import (Enricher, bucket_genres, catalog_candidates, clean_title, disc_number_hint,
                        pick_release)
from app.importer import import_workbook

JPEG = b"\xff\xd8\xff\xe0" + b"0" * 64

DIRT_US = "11111111-1111-1111-1111-111111111111"
DIRT_EU = "22222222-2222-2222-2222-222222222222"
DIRT_VINYL = "33333333-3333-3333-3333-333333333333"
DIRT_RG = "aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa"
TDS = "44444444-4444-4444-4444-444444444444"
TDS_RG = "bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb"


def release(mbid, title, artist, *, catno, fmt="CD", country="US", score=100, rg=DIRT_RG, date="1992-09-29",
            secondary=None):
    return {
        "id": mbid, "score": score, "title": title, "status": "Official", "date": date, "country": country,
        "artist-credit": [{"name": artist, "artist": {"name": artist}}],
        "release-group": {"id": rg, "primary-type": "Album", "secondary-types": secondary or []},
        "label-info": [{"catalog-number": catno, "label": {"name": "Columbia"}}],
        "media": [{"format": fmt, "track-count": 13}],
    }


SEARCH = {
    "Dirt": [
        release(DIRT_VINYL, "Dirt", "Alice in Chains", catno="C 52475", fmt="12\" Vinyl", score=100),
        release(DIRT_EU, "Dirt", "Alice in Chains", catno="472330 2", country="XE", score=100),
        release(DIRT_US, "Dirt", "Alice in Chains", catno="CK 52475", score=97),
    ],
    "The Downward Spiral": [
        release(TDS, "The Downward Spiral", "Nine Inch Nails", catno="INTD-92346", rg=TDS_RG, date="1994-03-08"),
    ],
    "Nevermind": [
        release("55555555-5555-5555-5555-555555555555", "Nevermind It's an Interview", "Nirvana", catno="PRO-CD-4429"),
    ],
}

DIRT_TRACKS = ["Them Bones", "Dam That River", "Rain When I Die", "Down in a Hole", "Sickman", "Rooster",
               "Junkhead", "Dirt", "God Smack", "Untitled", "Hate to Feel", "Angry Chair", "Would?"]


def lookup(mbid):
    rg = TDS_RG if mbid == TDS else DIRT_RG
    first = "1994-03-08" if mbid == TDS else "1992-09-29"
    return {
        "id": mbid, "title": "x",
        "release-group": {"id": rg, "first-release-date": first},
        "label-info": [{"catalog-number": "CK 52475", "label": {"name": "Columbia"}}],
        "media": [{"position": 1, "format": "CD", "tracks": [
            {"position": i + 1, "number": str(i + 1), "title": t, "length": 200_000 + i * 1000}
            for i, t in enumerate(DIRT_TRACKS)]}],
    }


GENRES = {
    DIRT_RG: [{"name": "grunge", "count": 9}, {"name": "alternative metal", "count": 6}, {"name": "rock", "count": 2}],
    TDS_RG: [{"name": "industrial rock", "count": 8}, {"name": "industrial", "count": 5}],
}


APPLE_JPEG = b"\xff\xd8\xff\xe1" + b"A" * 64


def handler(calls, caa_down=False, mb_down=False):
    def handle(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        assert request.headers["User-Agent"].startswith("DiscLibrary/")
        path = request.url.path
        if request.url.host == "coverartarchive.org":
            if caa_down:
                raise httpx.ConnectTimeout("timed out")
            return httpx.Response(200, content=JPEG) if DIRT_RG in path else httpx.Response(404)
        if request.url.host == "itunes.apple.com":
            term = request.url.params["term"]
            if "Downward Spiral" in term or "Dirt" in term:
                name = "The Downward Spiral" if "Downward" in term else "Dirt"
                artist = "Nine Inch Nails" if "Downward" in term else "Alice In Chains"
                return httpx.Response(200, json={"results": [
                    {"collectionName": "Something Else", "artistName": "Other", "artworkUrl100": "https://is1-ssl.mzstatic.com/x/100x100bb.jpg"},
                    {"collectionName": name, "artistName": artist, "artworkUrl100": "https://is1-ssl.mzstatic.com/a/b/100x100bb.jpg"}]})
            return httpx.Response(200, json={"results": []})
        if request.url.host == "is1-ssl.mzstatic.com":
            assert path.endswith("/1200x1200bb.jpg")
            return httpx.Response(200, content=APPLE_JPEG)
        if mb_down:
            raise httpx.ConnectTimeout("timed out")
        if path == "/ws/2/release/":
            q = request.url.params["query"]
            if q.startswith("barcode:"):
                return httpx.Response(200, json={"releases": [r for r in SEARCH["Dirt"] if r["id"] == DIRT_EU]})
            if q.startswith("(catno:"):
                return httpx.Response(200, json={"releases": SEARCH["Dirt"]})
            title = q.split('"')[1]
            return httpx.Response(200, json={"releases": SEARCH.get(title, [])})
        if path.startswith("/ws/2/release-group/"):
            return httpx.Response(200, json={"genres": GENRES.get(path.rsplit("/", 1)[1], [])})
        if path.startswith("/ws/2/release/"):
            mbid = path.rsplit("/", 1)[1]
            if mbid.startswith("9"):
                return httpx.Response(404, json={"error": "Not Found"})
            return httpx.Response(200, json=lookup(mbid))
        return httpx.Response(404)
    return handle


@pytest.fixture
def enricher(db, tmp_path):
    calls = []
    e = Enricher(db, tmp_path / "art", mb_interval=0,
                 client_factory=lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler(calls))))
    e.calls = calls
    return e


# --- cleaning and matching --------------------------------------------------------
def test_clean_title_and_disc_hint():
    assert clean_title("Live from Manchester — Yellow/Gold Disc") == "Live from Manchester"
    assert clean_title("Backstreet Boys — second copy (retail edition)") == "Backstreet Boys"
    assert clean_title("(What's the Story) Morning Glory?") == "(What's the Story) Morning Glory?"
    assert clean_title("Live in London (Blu-ray)") == "Live in London"
    assert disc_number_hint("Things We Do at Night — Disc 2") == 2
    assert disc_number_hint("Dirt") == 1
    assert disc_number_hint("Diamond Jubilee — Disc II") == 2
    assert disc_number_hint("101 — Disc B") == 2
    assert disc_number_hint("Workshop of the Telescopes (Disc One)") == 1
    assert disc_number_hint("All Eyez on Me — Book 2") == 2
    assert disc_number_hint("HIStory: Past, Present and Future, Book I — Disc 1: HIStory Begins") == 1
    assert clean_title("Tonight, Tonight (CD single)") == "Tonight, Tonight"
    assert clean_title("Root Down (EP)") == "Root Down"
    assert clean_title("Led Zeppelin IV (Untitled)") == "Led Zeppelin IV"
    assert clean_title("Counting Crows (The Best Of)") == "Counting Crows (The Best Of)"


def test_catalog_candidates():
    assert catalog_candidates("D 103309 (82596-2)") == ["D103309", "825962"]
    assert catalog_candidates("CK 52475") == ["CK52475"]
    for placeholder in (None, "—", "None", "Not printed on disc"):
        assert catalog_candidates(placeholder) == []


def test_catalog_number_beats_higher_search_score():
    best, why = pick_release(SEARCH["Dirt"], "Alice In Chains", "Dirt", "CK 52475")
    assert best["id"] == DIRT_US and why == "catalog number"


def test_without_catalog_prefers_cd_over_vinyl():
    best, _ = pick_release(SEARCH["Dirt"], "Alice In Chains", "Dirt", None)
    assert best["id"] in (DIRT_US, DIRT_EU)


def test_wrong_title_is_not_guessed():
    best, why = pick_release(SEARCH["Nevermind"], "Nirvana", "Nevermind", "DGCD-24425")
    assert best is None and ("mismatch" in why or "no matching" in why)


def test_bucket_genres():
    assert bucket_genres(GENRES[DIRT_RG]) == ("Grunge", "Grunge")
    assert bucket_genres(GENRES[TDS_RG]) == ("Industrial", "Industrial Rock")
    assert bucket_genres([{"name": "pop rock", "count": 3}]) == ("Pop", "Pop Rock")
    assert bucket_genres([]) == (None, None)


# --- the job ----------------------------------------------------------------------
def test_job_enriches_library(db, workbook, enricher):
    import_workbook(db, workbook)

    async def go():
        enricher.start()
        await enricher.wait()
    asyncio.run(go())

    dirt = db.one("SELECT * FROM albums WHERE title = 'Dirt'")
    assert (dirt["mbid_release"], dirt["enrich_status"], dirt["genre"], dirt["subgenre"]) == (DIRT_US, "matched", "Grunge", "Grunge")
    assert dirt["original_year"] == 1992 and dirt["art_url"].startswith(f"art/{dirt['id']}.jpg")
    assert (enricher.art_dir / f"{dirt['id']}.jpg").read_bytes() == JPEG
    tracks = db.query("SELECT number, title, duration_sec FROM tracks WHERE album_id = ? ORDER BY number", (dirt["id"],))
    assert len(tracks) == 13 and tracks[-1]["title"] == "Would?"
    assert dirt["runtime_sec"] == sum(t["duration_sec"] for t in tracks)

    tds = db.one("SELECT * FROM albums WHERE title = 'The Downward Spiral'")
    assert tds["genre"] == "Industrial" and tds["art_source"] == "apple"   # no CAA art -> Apple fallback
    assert (enricher.art_dir / f"{tds['id']}.jpg").read_bytes() == APPLE_JPEG
    assert dirt["art_source"] == "caa" and dirt["match_method"] == "catalog" and tds["match_method"] == "catalog"
    assert (dirt["rel_catno"], dirt["rel_country"], dirt["rel_format"]) == ("CK 52475", None, "CD")

    nevermind = db.one("SELECT * FROM albums WHERE title = 'Nevermind'")
    assert nevermind["enrich_status"] == "not_found"
    assert db.one("SELECT enrich_status FROM albums WHERE media_kind = 'video' LIMIT 1")["enrich_status"] == "pending"

    s = enricher.state
    assert (s.running, s.done, s.matched) == (False, s.total, 2)
    hosts = {httpx.URL(c).host for c in enricher.calls}
    assert hosts <= {"musicbrainz.org", "coverartarchive.org", "itunes.apple.com", "is1-ssl.mzstatic.com"}


def test_user_genre_is_kept(db, workbook, enricher):
    import_workbook(db, workbook)
    db.execute("UPDATE albums SET genre = 'Seattle' WHERE title = 'Dirt'")
    aid = db.one("SELECT id FROM albums WHERE title = 'Dirt'")["id"]
    asyncio.run(enricher.enrich_one(aid))
    assert db.one("SELECT genre FROM albums WHERE id = ?", (aid,))["genre"] == "Seattle"


def test_manual_fix_with_link(db, workbook, enricher):
    import_workbook(db, workbook)
    aid = db.one("SELECT id FROM albums WHERE title = 'Nevermind'")["id"]
    r = asyncio.run(enricher.enrich_one(aid, f"https://musicbrainz.org/release/{DIRT_EU}"))
    assert r["status"] == "manual"
    assert db.one("SELECT mbid_release FROM albums WHERE id = ?", (aid,))["mbid_release"] == DIRT_EU
    bad = asyncio.run(enricher.enrich_one(aid, "not a link"))
    assert bad["status"] == "error"
    gone = asyncio.run(enricher.enrich_one(aid, "99999999-9999-9999-9999-999999999999"))
    assert gone["status"] == "not_found"


def test_migration_adds_columns_to_old_db(tmp_path):
    import sqlite3
    from app.db import SCHEMA, Database
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.executescript(SCHEMA)
    con.execute("INSERT INTO albums (import_key, artist, title) VALUES ('k', 'A', 'T')")
    con.commit()
    con.close()
    db = Database(path)
    row = db.one("SELECT enrich_status, mbid_release FROM albums")
    assert row == {"enrich_status": "pending", "mbid_release": None}


def make_enricher(db, tmp_path, **kw):
    calls = []
    e = Enricher(db, tmp_path / "art2", mb_interval=0,
                 client_factory=lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler(calls, **kw))))
    e.calls = calls
    return e


def run_job(e):
    async def go():
        e.start()
        await e.wait()
    asyncio.run(go())


def test_unreachable_cover_archive_switches_to_apple(db, workbook, tmp_path):
    import_workbook(db, workbook)
    e = make_enricher(db, tmp_path, caa_down=True)
    run_job(e)
    dirt = db.one("SELECT * FROM albums WHERE title = 'Dirt'")
    assert dirt["art_source"] == "apple" and dirt["enrich_status"] == "matched"
    caa_calls = [c for c in e.calls if "coverartarchive" in c]
    assert len(caa_calls) <= 3   # gives up on the archive after a few failures instead of timing out on every album


def test_job_stops_when_musicbrainz_unreachable(db, workbook, tmp_path):
    import_workbook(db, workbook)
    e = make_enricher(db, tmp_path, mb_down=True)
    run_job(e)
    assert e.state.done == 3 and "can't be reached" in e.state.stopped_reason
    assert db.one("SELECT COUNT(*) AS n FROM albums WHERE enrich_status = 'pending'")["n"] >= 1


def test_search_pressings_by_catalog_and_barcode(db, workbook, enricher):
    import_workbook(db, workbook)
    aid = db.one("SELECT id FROM albums WHERE title = 'Dirt'")["id"]
    by_cat = asyncio.run(enricher.search_pressings(aid, "CK 52475", None))
    assert by_cat[0]["id"] == DIRT_US and by_cat[0]["catalog_match"] and by_cat[0]["format"] == "CD"
    assert by_cat[-1]["format"] == '12" Vinyl'                      # vinyl sorts last
    by_bc = asyncio.run(enricher.search_pressings(aid, None, "0 74645-24752 6"))
    assert [c["id"] for c in by_bc] == [DIRT_EU]
    assert any(httpx.URL(c).params.get("query") == "barcode:074645247526" for c in enricher.calls)


def test_lock_pins_release_and_job_leaves_it_alone(db, workbook, enricher):
    import_workbook(db, workbook)
    aid = db.one("SELECT id FROM albums WHERE title = 'Dirt'")["id"]
    r = asyncio.run(enricher.lock(aid, f"https://musicbrainz.org/release/{DIRT_EU}", "472330 2", None))
    assert r["status"] == "locked"
    a = db.one("SELECT * FROM albums WHERE id = ?", (aid,))
    assert (a["pressing_locked"], a["mbid_release"], a["catalog_verified"], a["catalog_no"]) == (1, DIRT_EU, "472330 2", "CK 52475")
    db.execute("UPDATE albums SET enrich_status = 'pending' WHERE id = ?", (aid,))
    run_job(enricher)
    assert db.one("SELECT mbid_release FROM albums WHERE id = ?", (aid,))["mbid_release"] == DIRT_EU
    enricher.unlock(aid)
    assert db.one("SELECT pressing_locked FROM albums WHERE id = ?", (aid,))["pressing_locked"] == 0
    bad = asyncio.run(enricher.lock(aid, "nope", None, None))
    assert bad["status"] == "error"


def test_reimport_never_touches_verified_pressing(db, workbook, enricher):
    import_workbook(db, workbook)
    aid = db.one("SELECT id FROM albums WHERE title = 'Dirt'")["id"]
    asyncio.run(enricher.lock(aid, DIRT_US, None, None))
    import_workbook(db, workbook)
    a = db.one("SELECT pressing_locked, catalog_verified, mbid_release FROM albums WHERE id = ?", (aid,))
    assert a == {"pressing_locked": 1, "catalog_verified": "CK 52475", "mbid_release": DIRT_US}


def test_pressing_api(db, workbook, enricher, tmp_path, drivers):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.config import ChangerConfig, Settings
    from app.ha_client import HAClient
    from app.main import create_app
    www = tmp_path / "www"
    www.mkdir()
    (www / "index.html").write_text("ok")
    settings = Settings(data_dir=tmp_path, www_dir=www, changers=[ChangerConfig(1, "C1"), ChangerConfig(2, "C2")])
    from app.db import Database
    app = create_app(settings, drivers=drivers, ha=HAClient(token=""), enricher=Enricher(
        Database(settings.db_path), tmp_path / "art3", mb_interval=0, client_factory=enricher.client_factory))
    with TestClient(app) as c:
        c.post("/api/import", files={"file": ("x.xlsx", workbook.read_bytes())})
        dirt = next(a for a in c.get("/api/albums").json() if a["title"] == "Dirt")
        r = c.post(f"/api/albums/{dirt['id']}/enrich", json={})
        assert r.json()["result"]["status"] == "matched"
        p = c.get("/api/pressings?filter=catalog").json()
        assert p["counts"]["catalog"] >= 1 and any(a["title"] == "Dirt" for a in p["albums"])
        cands = c.get(f"/api/albums/{dirt['id']}/pressings", params={"catalog_no": "CK 52475"}).json()
        assert cands[0]["is_current"] and cands[0]["catalog_match"]
        assert c.post("/api/pressings/lock-catalog-matches").json()["locked"] >= 1
        assert c.get("/api/pressings?filter=locked").json()["counts"]["locked"] >= 1
        assert c.post(f"/api/albums/{dirt['id']}/lock", json={"mbid": "garbage"}).status_code == 409
        x = c.get("/api/pressings/export")
        assert x.status_code == 200
        out = tmp_path / "exp.xlsx"
        out.write_bytes(x.content)
        ws = load_workbook(out).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "Artist" and any(r[1] == "Dirt" and r[5] == "Yes" for r in rows)
