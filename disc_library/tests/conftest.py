import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import ChangerConfig  # noqa: E402
from app.db import Database  # noqa: E402
from app.drivers.simulated import SimulatedChanger  # noqa: E402


class FakeClock:
    def __init__(self):
        self.t = 1000.0

    def __call__(self):
        return self.t

    def advance(self, seconds):
        self.t += seconds


@pytest.fixture
def clock():
    return FakeClock()


@pytest.fixture
def db(tmp_path):
    return Database(tmp_path / "test.db")


@pytest.fixture
def drivers(clock):
    return {
        1: SimulatedChanger(ChangerConfig(1, "Changer 1"), track_seconds=10, clock=clock),
        2: SimulatedChanger(ChangerConfig(2, "Changer 2"), track_seconds=10, clock=clock),
    }


@pytest.fixture
def workbook(tmp_path):
    """Mirrors the real catalog: three tabs, '#' column, statuses, a title row above the header."""
    wb = Workbook()
    ws = wb.active
    ws.title = "CD Collection"
    ws.append(["#", "Artist", "Album", "CD ID / Catalog No.", "Year", "Label", "Edition / Notes", "Source", "Status"])
    ws.append([1, "Nine Inch Nails", "The Downward Spiral", "INTD-92346", 1994, "Nothing/Interscope", "", "", "In collection"])
    ws.append([2, "Alice In Chains", "Dirt", "CK 52475", "1992", "Columbia", "", "", "In collection"])
    ws.append([3, "Nirvana", "Nevermind", "DGCD-24425", 1991, "DGC", "BMG Direct club edition", "", "In collection"])
    ws.append([4, "Radiohead", "Kid A", "", 2000, "Capitol", "", "", "On order"])
    ws.append([None, None, None, None, None, None, None, None, None])
    ws2 = wb.create_sheet("Non-CD Formats")
    ws2.append(["Artist", "Title", "Format", "Catalog No.", "Year", "Label", "Notes", "Status"])
    ws2.append(["Pink Floyd", "The Dark Side of the Moon", "SACD", "", 1973, "Harvest", "", "Owned"])
    ws2.append(["Example Band", "Live Somewhere", "Blu-ray (concert video)", "", 2015, "", "", "Owned"])
    ws3 = wb.create_sheet("Video (non-CD)")
    ws3.insert_rows(1)
    ws3["A1"] = "Video collection"
    ws3.append(["Artist", "Album", "Format", "Year", "Status"])
    ws3.append(["Various", "Example Concert Film", "Blu-ray", 2010, "Owned"])
    path = tmp_path / "CD_Collection.xlsx"
    wb.save(path)
    return path
